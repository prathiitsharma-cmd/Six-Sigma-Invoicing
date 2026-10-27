"""Stage 4, step 2 - rigorous uncertainty for baseline / audit DPMO and the reduction: exact Poisson, overdispersion (quasi-Poisson),
day-level cluster bootstrap, NB regression; Monte Carlo intervals + standard errors for Hybrid H.
Reads : data/original.xlsx, outputs/results.json (or precomputed/), 03_stress/comp.py (exec-ed)
Writes: outputs/uncert.json
"""
# ---------------------------------------------------------------------------
# WHAT THIS FILE DOES (plain English)
#   Gives honest error bars on the headline numbers: exact Poisson and over-dispersion-adjusted intervals, a bootstrap that
#   resamples whole audit days (B = 20,000, seed 2026), a Bayesian check, and a 20,000-run Monte Carlo (seed 5) for Hybrid H.
# ---------------------------------------------------------------------------
import sys, pathlib
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))   # packaging glue: lets us "import paths"
import paths                                                            # packaging glue: every file location lives in paths.py
import json, math, numpy as np
from openpyxl import load_workbook
from scipy.stats import chi2, norm, gamma
wb = load_workbook(paths.SRC, data_only=True)   # identical file to the original upload (byte-for-byte)
def rows(n): return [r[:4] for r in wb[n].iter_rows(min_row=3, values_only=True) if r[0] is not None and hasattr(r[0], 'year')]
M = rows('Measure'); I = rows('Improve')
Madj = [r for r in M if not (r[0].month == 8 and r[0].day == 25)]
OPP = 58
def dpmo(e, n): return e / n / OPP * 1e6
def exact(e, n):
    lo = chi2.ppf(0.025, 2 * e) / 2 if e > 0 else 0.0; hi = chi2.ppf(0.975, 2 * e + 2) / 2
    return dpmo(lo, n), dpmo(hi, n)
def disp(rs):
    e = np.array([r[3] for r in rs], float); n = np.array([r[2] for r in rs], float); u = e.sum() / n.sum()
    return float(((e - n * u) ** 2 / (n * u)).sum() / (len(rs) - 1))      # Pearson dispersion (Poisson = 1)
eb, nb = sum(r[3] for r in Madj), sum(r[2] for r in Madj); ei, ni = sum(r[3] for r in I), sum(r[2] for r in I)
out = {}
out['base'] = dict(errors=eb, audited=nb, dpmo=dpmo(eb, nb), exact=exact(eb, nb), phi=disp(Madj))
out['imp'] = dict(errors=ei, audited=ni, dpmo=dpmo(ei, ni), exact=exact(ei, ni), phi=disp(I))
for k in ('base', 'imp'):
    o = out[k]; se = math.sqrt(o['phi'] * o['errors']); o['quasi'] = (dpmo(o['errors'] - 1.96 * se, o[k == 'base' and 'audited' or 'audited']), dpmo(o['errors'] + 1.96 * se, o['audited']))
# day-level cluster bootstrap
rng = np.random.default_rng(2026); B = 20000
eB = np.array([r[3] for r in Madj], float); nB = np.array([r[2] for r in Madj], float)
eI = np.array([r[3] for r in I], float); nI = np.array([r[2] for r in I], float)
ib = rng.integers(0, len(eB), (B, len(eB))); ii = rng.integers(0, len(eI), (B, len(eI)))
db = eB[ib].sum(1) / nB[ib].sum(1) / OPP * 1e6; di = eI[ii].sum(1) / nI[ii].sum(1) / OPP * 1e6
red = 1 - di / db
out['boot'] = dict(B=B, base=(float(np.percentile(db, 2.5)), float(np.percentile(db, 97.5))), imp=(float(np.percentile(di, 2.5)), float(np.percentile(di, 97.5))),
                   red=(float(np.percentile(red, 2.5)), float(np.percentile(red, 97.5))), p_imp_gt500=float(np.mean(di > 500)), red_lo95_one_sided=float(np.percentile(red, 5)))
# Bayesian (Jeffreys-type gamma posterior on the audit error rate)
lam500 = 500 * ni * OPP / 1e6
out['bayes'] = dict(p_gt500=float(1 - gamma.cdf(lam500, ei + 1)), p_gt500_jeffreys=float(1 - gamma.cdf(lam500, ei + 0.5)))
# NB regression (overdispersion- and peak-adjusted) reduction
R = json.load(open(paths.find('results.json')))['regression']['coefs']['post']
out['nb'] = dict(irr=R['irr'], lo=R['lo'], hi=R['hi'], red=(1 - R['hi'], 1 - R['lo']), red_point=1 - R['irr'])
# Monte Carlo intervals for Hybrid H (20,000 runs, seed 5 as in stage3)
exec(open(paths.STRESS / 'comp.py').read())
Hm = mk(['A1','A2','A3','A4','B1','B2','B3','B4','B5','C2']); Pm = mc_draws(20000, seed=5)
r = evaluate(Hm, Pm); N = len(r['cut'])
pc = lambda a, q: float(np.percentile(a, q))
out['mcH'] = dict(cut=(pc(r['cut'], 5), pc(r['cut'], 50), pc(r['cut'], 95)), dpmo=(pc(r['dpmo'], 5), pc(r['dpmo'], 50), pc(r['dpmo'], 95)),
                  npv=(pc(r['npv'], 5), pc(r['npv'], 50), pc(r['npv'], 95)), pneg=float(np.mean(r['npv'] < 0)), p75=float(np.mean(r['cut'] >= .75)), N=N)
for k in ('pneg', 'p75'):
    p = out['mcH'][k]; out['mcH'][k + '_se'] = math.sqrt(p * (1 - p) / N)
json.dump(out, open(paths.out('uncert.json'), 'w'), indent=1, default=float)
for k, v in out.items(): print(k, v)
