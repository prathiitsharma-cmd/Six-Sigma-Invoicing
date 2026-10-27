"""Stage 5 - build the Excel model (model2.xlsx) from the data and all earlier results.
Reads : data/original.xlsx, sol.json, parts.json, stage3.json, marginal.json, results.json, uncert.json, rule.json,
        rule_final.json, laney.json (each from outputs/ if present, else precomputed/)
Writes: outputs/model2.xlsx (formulas, no cached numbers), outputs/uref.json (map of named result cells)
"""
# ---------------------------------------------------------------------------
# WHAT THIS FILE DOES (plain English)
#   Writes every sheet of the workbook with openpyxl: inputs, baseline, solutions A/B/C, comparison, plus (by running the four
#   ext_*.py files in this folder) the Hybrid, Parts Library, Stress Test, Uncertainty, Control Chart and Reconciliation sheets.
#   Most cells are live Excel formulas; Monte Carlo / bootstrap results are pasted in as numbers. fullCalcOnLoad is set, so
#   Excel computes every formula when you open the file. recalc_and_inject.py is OPTIONAL (only stores cached values).
# ---------------------------------------------------------------------------
import sys, pathlib
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))   # packaging glue: lets us "import paths"
import paths                                                            # packaging glue: every file location lives in paths.py
import json
import pandas as pd, numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import CellIsRule, DataBarRule

SRC = paths.SRC
S = json.load(open(paths.find('sol.json')))
wb = openpyxl.Workbook()
FONT = 'Arial'
NAVY = '13294B'; LBLUE = 'DCE8F5'; GREY = '595959'
F_HDR = PatternFill('solid', fgColor=NAVY); F_IN = PatternFill('solid', fgColor='FFF2CC'); F_SUB = PatternFill('solid', fgColor=LBLUE)
F_KPI = PatternFill('solid', fgColor='EEF3F9'); F_TOT = PatternFill('solid', fgColor='F2F2F2')
thin = Side(style='thin', color='BFBFBF'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BLUE, BLACK, GREEN = '0000FF', '000000', '008000'

def f(sz=10, b=False, c=BLACK, i=False): return Font(name=FONT, size=sz, bold=b, color=c, italic=i)
def put(ws, ref, v, fmt=None, b=False, c=BLACK, fill=None, al=None, wrap=False, sz=10, border=True, i=False):
    cell = ws[ref]; cell.value = v; cell.font = f(sz, b, c, i)
    if fmt: cell.number_format = fmt
    if fill: cell.fill = fill
    if border: cell.border = BOX
    if al or wrap: cell.alignment = Alignment(horizontal=al, vertical='center', wrap_text=wrap)
    return cell
def inp(ws, ref, v, fmt=None): return put(ws, ref, v, fmt, c=BLUE, fill=F_IN, al='center')
def link(ws, ref, v, fmt=None, b=False): return put(ws, ref, v, fmt, c=GREEN, al='center', b=b)
def calc(ws, ref, v, fmt=None, b=False, fill=None): return put(ws, ref, v, fmt, b=b, al='center', fill=fill)
def hdr(ws, r, c0, labels, h=30):
    for i, l in enumerate(labels):
        cell = ws.cell(row=r, column=c0 + i, value=l); cell.font = f(9, True, 'FFFFFF'); cell.fill = F_HDR
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cell.border = BOX
    ws.row_dimensions[r].height = h
def title(ws, t, sub):
    ws['A1'] = t; ws['A1'].font = f(16, True, NAVY)
    ws['A2'] = sub; ws['A2'].font = f(9, False, GREY, True)
    ws.sheet_view.showGridLines = False
def section(ws, ref, t):
    ws[ref] = t; ws[ref].font = f(12, True, NAVY)
def name(n, ref): wb.defined_names[n] = DefinedName(n, attr_text=ref)
def widths(ws, d):
    for k, v in d.items(): ws.column_dimensions[k].width = v

# ============================================================ README
RD = wb.active; RD.title = 'README'
title(RD, 'Invoice-Error Solutions Model · Break The Case \'26', 'Three solution options for the Green Energy GCC invoice-error case, modelled on the case workbook. Base-case results are live formulas; stress-test, search, bootstrap and regression results are pasted from the Python model and labelled as such.')
widths(RD, {'A': 30, 'B': 100})
rows = [('Purpose', 'Compare three solutions (A Digital Front Door, B Lean Flow & Peak Control, C Supplier Quality Partnership) and their combination on error reduction, capacity and money.'),
        ('How to use', '1) Pick a scenario on Inputs (Central / Conservative / Optimistic). 2) Change any yellow cell with blue text. 3) Read results on each solution sheet and on Comparison.'),
        ('Sheets', 'Inputs · Baseline (case data) · Sol A · Sol B · Sol C · Combined (A+B+C) · Comparison (decision matrix + charts) · Sensitivity · Assumptions'),
        ('Colour key', 'Blue text on yellow = input you can change · Black = formula · Green = linked from another sheet'),
        ('Error model', 'Each solution removes a share of each of the 19 baseline error reasons (effectiveness input). Errors removed = reason share × effectiveness. Month-end "peak excess" errors are cut separately. Total cut = 1 − (1 − reason cut) × (1 − peak cut).'),
        ('Money model', 'Benefit = baseline cost of poor quality (GST penalties + exception rework, ₹ lakh/yr) × total cut, plus any extra operating saving. Headline NPV and payback count build time: each part pays from the month after it goes live. Simple versions (benefits from day 1) are kept and labelled.'),
        ('Units', '₹ L = ₹ lakh (₹1,00,000). DPMO = defects per million opportunities (58 fields per audited invoice).'),
        ('Limits', 'Solution costs and effectiveness values are illustrative analyst assumptions (the case gives none). Simulation outputs (waiting days, backlog) come from a separate Python model and are pasted as values.'),
        ('Targets', 'Pass mark = ≥ 75% cut in errors per invoice. ≤ 500 DPMO is a team-set stretch goal (it is not in the case brief); no mix of parts reaches it in the model, so Hybrid H carries a gate that escalates to H+.')]
for i, (a, b) in enumerate(rows):
    put(RD, f'A{4+i}', a, b=True, fill=F_SUB); put(RD, f'B{4+i}', b, wrap=True); RD.row_dimensions[4 + i].height = 32

# ============================================================ INPUTS
IN = wb.create_sheet('Inputs')
title(IN, 'Inputs', 'Global assumptions. Yellow cells with blue text are editable.')
widths(IN, {'A': 52, 'B': 14, 'C': 16, 'D': 70})
hdr(IN, 4, 1, ['Parameter', 'Value', 'Unit', 'Source / note'])
items = [('WD', 'Working days per month', 26, '0', 'days', 'Case Process Steps tab'),
         ('OPP', 'Opportunities per audited invoice', 58, '0', 'fields', 'Case Glossary'),
         ('PEN_AMT', 'GST penalty per affected invoice', 5000, '"₹"#,##0', 'INR', 'Case statement (up to ₹5,000)'),
         ('PEN_BASE', 'Penalty trigger rate (central)', 0.05, '0.0%', '% of erroneous invoices', 'Case Cost Assumptions'),
         ('DEF_SHARE', 'Share of invoices with ≥1 error', 0.35, '0.0%', '% of invoices', 'Case Cost Assumptions (data implies ~46%)'),
         ('EXC_SHARE', 'Share of invoices reaching exception handling', 0.25, '0.0%', '% of invoices', 'Case Process Steps step 7'),
         ('EXC_MIN', 'Exception touch time', 12, '0', 'minutes', 'Case Process Steps step 7'),
         ('RATE', 'Analyst fully loaded cost', 350, '"₹"#,##0', 'INR/hour', 'Case Cost Assumptions'),
         ('MONTHS', 'Months per year', 12, '0', 'months', ''),
         ('DISC', 'Discount rate', 0.12, '0.0%', '% per year', 'Analyst assumption'),
         ('BASE_TOUCH', 'Baseline weighted touch time per invoice', 23, '0.0', 'minutes', 'Case Process Steps (steps 1–9, weighted)')]
r = 5
for nm, lab, v, fmt, unit, src in items:
    put(IN, f'A{r}', lab); inp(IN, f'B{r}', v, fmt); put(IN, f'C{r}', unit, c=GREY, sz=9); put(IN, f'D{r}', src, sz=9, c=GREY)
    name(nm, f"'Inputs'!$B${r}"); r += 1
r += 1
section(IN, f'A{r}', 'Scenario'); r += 1
put(IN, f'A{r}', 'Selected scenario', b=True); inp(IN, f'B{r}', 'Central'); name('SCEN', f"'Inputs'!$B${r}")
dv = DataValidation(type='list', formula1='"Central,Conservative,Optimistic"'); IN.add_data_validation(dv); dv.add(f'B{r}')
r += 2
hdr(IN, r, 1, ['Scenario', 'Effectiveness multiplier', 'Cost multiplier', 'Penalty trigger rate']); r += 1
sc0 = r
for nmv, a, b_, c_ in [('Central', 1, 1, '=PEN_BASE'), ('Conservative', 0.6, 1.5, 0.02), ('Optimistic', 1.15, 0.8, 0.08)]:
    put(IN, f'A{r}', nmv, b=True); inp(IN, f'B{r}', a, '0.00'); inp(IN, f'C{r}', b_, '0.00'); inp(IN, f'D{r}', c_, '0.0%'); r += 1
r += 1
put(IN, f'A{r}', 'Active effectiveness multiplier', b=True); calc(IN, f'B{r}', f'=INDEX(B{sc0}:B{sc0+2},MATCH(SCEN,A{sc0}:A{sc0+2},0))', '0.00', True); name('M_EFF', f"'Inputs'!$B${r}"); r += 1
put(IN, f'A{r}', 'Active cost multiplier', b=True); calc(IN, f'B{r}', f'=INDEX(C{sc0}:C{sc0+2},MATCH(SCEN,A{sc0}:A{sc0+2},0))', '0.00', True); name('M_COST', f"'Inputs'!$B${r}"); r += 1
put(IN, f'A{r}', 'Active penalty trigger rate', b=True); calc(IN, f'B{r}', f'=INDEX(D{sc0}:D{sc0+2},MATCH(SCEN,A{sc0}:A{sc0+2},0))', '0.0%', True); name('PEN', f"'Inputs'!$B${r}"); r += 2
section(IN, f'A{r}', 'Derived'); r += 1
put(IN, f'A{r}', 'Monthly invoice volume (baseline avg/day × working days)'); link(IN, f'B{r}', "='Baseline'!B7*WD", '#,##0'); name('VOL', f"'Inputs'!$B${r}"); r += 1
put(IN, f'A{r}', 'Baseline COPQ – GST penalties (₹ L/yr)'); calc(IN, f'B{r}', '=VOL*MONTHS*DEF_SHARE*PEN*PEN_AMT/100000', '0.0'); name('COPQ_PEN', f"'Inputs'!$B${r}"); r += 1
put(IN, f'A{r}', 'Baseline COPQ – exception rework (₹ L/yr)'); calc(IN, f'B{r}', '=VOL*MONTHS*EXC_SHARE*EXC_MIN/60*RATE/100000', '0.0'); name('COPQ_RW', f"'Inputs'!$B${r}"); r += 1
put(IN, f'A{r}', 'Baseline COPQ – total (₹ L/yr)', b=True); calc(IN, f'B{r}', '=COPQ_PEN+COPQ_RW', '0.0', True, F_KPI); name('COPQ', f"'Inputs'!$B${r}"); r += 1
IN.freeze_panes = 'A5'

# ============================================================ BASELINE
BS = wb.create_sheet('Baseline')
title(BS, 'Baseline (case data)', 'Measure tab Aug–Sep 2026 and Improve tab Mar–Apr 2027, plus the Error Reasons Pareto. Adjusted baseline excludes 25-Aug (0 errors on a 659-invoice peak day).')
widths(BS, {'A': 46, 'B': 14, 'C': 4, 'D': 13, 'E': 11, 'F': 11, 'G': 11, 'H': 9, 'I': 11, 'J': 4, 'K': 46, 'L': 11, 'M': 11})
M = pd.read_excel(SRC, sheet_name='Measure', header=1, nrows=52); M.columns = ['Date', 'Inv', 'Samp', 'Err']
I = pd.read_excel(SRC, sheet_name='Improve', header=1, nrows=51); I.columns = ['Date', 'Inv', 'Samp', 'Err']
hdr(BS, 4, 4, ['Date', 'Invoices', 'Samples', 'Errors', 'Peak? (≥400)', 'Include (1/0)'])
d0 = 5
for i, rr in M.iterrows():
    row = d0 + i
    put(BS, f'D{row}', rr.Date.to_pydatetime(), 'dd-mmm-yy', al='center', sz=9)
    for col, v in (('E', rr.Inv), ('F', rr.Samp), ('G', rr.Err)): put(BS, f'{col}{row}', int(v), '0', al='center', sz=9)
    calc(BS, f'H{row}', f'=IF(E{row}>=400,1,0)', '0'); calc(BS, f'I{row}', f'=IF(AND(G{row}=0,H{row}=1),0,1)', '0')
    BS[f'H{row}'].font = f(9); BS[f'I{row}'].font = f(9)
d1 = d0 + len(M) - 1
E_, S_, I_, P_, X_ = (f'{c}{d0}:{c}{d1}' for c in 'EFGHI')
section(BS, 'A4', 'Baseline summary')
summ = [('Audit days', f'=COUNT({E_})', '0'), ('Invoices received', f'=SUM({E_})', '#,##0'), ('Average invoices per day', f'=AVERAGE({E_})', '0.0'),
        ('Errors (adjusted)', f'=SUMIF({X_},1,{I_})', '#,##0'), ('Audited invoices (adjusted)', f'=SUMIF({X_},1,{S_})', '#,##0'),
        ('DPU – errors per audited invoice', '=B8/B9', '0.000'), ('DPMO', '=B10/OPP*1000000', '#,##0'), ('Sigma level', '=NORMSINV(1-B10/OPP)+1.5', '0.00'),
        ('Est. % invoices with ≥1 error (1 − e^−DPU)', '=1-EXP(-B10)', '0.0%'),
        ('DPU on peak days (adjusted)', f'=SUMIFS({I_},{P_},1,{X_},1)/SUMIFS({S_},{P_},1,{X_},1)', '0.000'),
        ('DPU on non-peak days', f'=SUMIFS({I_},{P_},0)/SUMIFS({S_},{P_},0)', '0.000'),
        ('Peak-excess errors (above non-peak rate)', f'=SUMIFS({I_},{P_},1,{X_},1)-SUMIFS({S_},{P_},1,{X_},1)*B15', '0.0'),
        ('Peak-excess share of all errors', '=B16/B8', '0.0%'),
        ('Improve audit: errors', int(I.Err.sum()), '0'), ('Improve audit: audited invoices', int(I.Samp.sum()), '#,##0'),
        ('Improve audit: DPMO (verified)', '=B18/B19/OPP*1000000', '#,##0'), ('Verified error reduction', '=1-(B18/B19)/B10', '0.0%')]
for k, (lab, fm, nf) in enumerate(summ):
    rr = 5 + k
    put(BS, f'A{rr}', lab, b=lab in ('DPMO', 'Sigma level', 'Verified error reduction'))
    if isinstance(fm, int): put(BS, f'B{rr}', fm, nf, al='center')
    else: calc(BS, f'B{rr}', fm, nf, lab in ('DPMO', 'Sigma level', 'Verified error reduction'), F_KPI if lab in ('DPMO', 'Sigma level', 'Verified error reduction') else None)
name('DPU0', "'Baseline'!$B$10"); name('DPMO0', "'Baseline'!$B$11"); name('EXCESS', "'Baseline'!$B$17"); name('VERIFIED', "'Baseline'!$B$21")

# ---- data-validation evidence, COPQ on two bases, capacity & triage (traceability for the deck)
section(BS, 'A23', 'Data validation evidence')
_ev = [('DPMO raw, incl. 25-Aug', f'=SUM({I_})/SUM({S_})/OPP*1000000', '#,##0'), ('Sigma level raw, incl. 25-Aug', f'=NORMSINV(1-SUM({I_})/SUM({S_})/OPP)+1.5', '0.00'),
       ('25-Aug: audited invoices', f'=SUMIFS({S_},D{d0}:D{d1},DATE(2026,8,25))', '0'), ('25-Aug: expected errors at peak-day DPU', '=B26*B14', '0.0'),
       ('25-Aug: expected errors at average DPU', '=B26*B10', '0.0'), ('25-Aug: P(0 errors) at average DPU (Poisson)', '=EXP(-B28)', '0.0E+00'),
       ('Supplied audit DPMO: 95% CI low (exact Poisson)', '=CHIINV(0.975,2*B18)/2/B19/OPP*1000000', '#,##0'), ('Supplied audit DPMO: 95% CI high (exact Poisson)', '=CHIINV(0.025,2*(B18+1))/2/B19/OPP*1000000', '#,##0'),
       ('P(true audit DPMO > 500)', '=1-GAMMADIST(500*B19*OPP/1000000,B18+1,1,TRUE)', '0%'), ('Est. % invoices with ≥1 error after (1 − e^−DPU)', '=1-EXP(-B18/B19)', '0.0%')]
for k, (lab, fm, nf) in enumerate(_ev):
    put(BS, f'A{24+k}', lab, sz=9); calc(BS, f'B{24+k}', fm, nf)
section(BS, 'A35', 'Cost of poor quality on two bases (₹ L/yr)')
_cq = [('COPQ after, case basis (35%) = COPQ × (1 − verified cut)', '=COPQ*(1-B21)', '0.0'), ('COPQ baseline, Poisson basis (≈46% with errors)', '=VOL*MONTHS*B13*PEN*PEN_AMT/100000+COPQ_RW', '0.0'),
       ('COPQ after, Poisson basis', '=VOL*MONTHS*B33*PEN*PEN_AMT/100000+COPQ_RW*(1-B21)', '0.0')]
for k, (lab, fm, nf) in enumerate(_cq):
    put(BS, f'A{36+k}', lab, sz=9); calc(BS, f'B{36+k}', fm, nf)
section(BS, 'A40', 'Capacity and triage load')
_cp = [('Max invoices, normal day (18 analysts × 8 h ÷ touch time)', '=18*8*60/BASE_TOUCH', '#,##0'), ('Max invoices, peak day (+3 h overtime)', '=18*11*60/BASE_TOUCH', '#,##0'),
       ('Average load on normal hours', '=B7/B41', '0%'), ('Monthly load incl. peak-day overtime (20 normal + 6 peak days)', '=B7*WD/((WD-6)*B41+6*B42)', '0.0%'),
       ('Triage hours, average day (2 min per invoice)', '=B7*2/60', '0.0'), ('Triage hours, worst day', f'=MAX({E_})*2/60', '0.0'),
       ('Triage hours, worst day after the calendar (25% of volume above 400 moved)', f'=(MAX({E_})-0.25*(MAX({E_})-400))*2/60', '0.0')]
for k, (lab, fm, nf) in enumerate(_cp):
    put(BS, f'A{41+k}', lab, sz=9); calc(BS, f'B{41+k}', fm, nf)
# Pareto
ER = pd.read_excel(SRC, sheet_name='Error Reasons', nrows=19); ER.columns = ['cat', 'reason', 'n', 'src']
ER = ER.sort_values('n', ascending=False, kind='stable').reset_index(drop=True)
hdr(BS, 4, 11, ['Error reason (Pareto order)', 'Errors', 'Share']); BS['K4'].alignment = Alignment(horizontal='left', vertical='center')
for k, rr in ER.iterrows():
    row = 5 + k
    put(BS, f'K{row}', rr.reason, sz=9); put(BS, f'L{row}', int(rr.n), '0', al='center', sz=9)
    calc(BS, f'M{row}', f'=L{row}/SUM($L$5:$L$23)', '0.0%')
put(BS, 'K24', 'Total', b=True, fill=F_TOT); calc(BS, 'L24', '=SUM(L5:L23)', '0', True, F_TOT); calc(BS, 'M24', '=SUM(M5:M23)', '0%', True, F_TOT)
BS.freeze_panes = 'A5'
REASONS = list(ER.reason)

# ============================================================ SOLUTION SHEETS
STEPS = ['1 Receive invoice', '2 Log into ERP', '3 Format & field validation', '4 Vendor verification', '5 PO matching', '6 Approval routing', '8 GST & tax validation', '9 Final approval wait', 'Triage check (Solution B only)']
BASE_STEPS = [1, 1, 5, 3, 5, 2, 3, 0, 0]
SOLS = {
 'A': dict(sheet='Sol A – Digital Front Door', name='Digital Front Door', tag='Stop bad invoices before they enter (technology-led prevention at submission).',
           eff=S['eff']['S1'], peak=S['peakfix']['S1'], steps=[0.5, 0.2, 1, 0.5, 1.5, 1, 1, 0, 0], months=5, extra=0,
           capex=[('Supplier portal (ERP-native, licences + setup)', 25), ('Rules engine R1–R14, IRN feed, ERP integration', 12), ('Approval workflow + SLA timers', 8)],
           opex=[('Portal licences & support', 6), ('Integration monitoring & rule upkeep', 4)],
           sim=dict(late=S['sol']['S1']['late2'], maxbl=S['sol']['S1']['maxbl']),
           comps=['Supplier portal + e-invoice (IRN) feed replace email intake', 'Rules R1–R14 at submission: 3-way PO/GRN match ±1%, PO-line picking, mandatory delivery note, duplicate block, vendor-master fields, currency lock, auto GST',
                  'Valid invoices auto-create the ERP record (no re-keying)', 'Rules-based approval routing with 24 h SLA', 'Human-in-the-loop queue for soft fails / risk ≥ 22%', 'Live dashboard: rejections by vendor, daily u-chart'],
           risks=[('Small vendors cannot use a portal', 'Assisted submission desk; mobile upload; phase by segment'), ('Rules too strict → good invoices rejected', 'Tune tolerances in pilot (hold if P(error) > 22%)'),
                  ('IT build slips', 'Buy ERP-native; launch rules R1–R4 first (77% of errors)'), ('Vendor resistance', 'Procurement mandates one channel; scorecards')]),
 'B': dict(sheet='Sol B – Lean Flow', name='Lean Flow & Peak Control', tag='Standardise the work and flatten month-end with what exists today (process & people).',
           eff=S['eff']['S2'], peak=S['peakfix']['S2'], steps=[1, 1, 5, 3, 5, 2, 3, 0, 2], months=1.5, extra=round(S['sol']['S2']['benefit'] - S['sol']['S2']['cut'] * 104.9769, 2),
           capex=[('Templates, checklists, vendor communication', 2), ('Training & standard work', 2), ('Dashboard / u-chart setup', 2)],
           opex=[('Flex-pool coordination & training refresh', 3), ('Kaizen & dashboard upkeep', 2)],
           sim=dict(late=S['sol']['S2']['late2'], maxbl=S['sol']['S2']['maxbl']),
           comps=['Intake triage desk: 2-minute completeness checklist; same-day bounce with reason', 'Standard invoice template + submission checklist for all vendors', 'Month-end submission calendar (≈25% of peak volume moved earlier)',
                  '4-person flex pool on the last 7 working days + 3 h overtime', 'Approval SLA 24 h with daily escalation list', 'Standard work, daily u-chart huddle, monthly Kaizen'],
           risks=[('Gains plateau near 50%', 'Use as phase 1, then automate'), ('Checklist discipline slips at month-end', 'Daily huddle; flex pool covers triage'),
                  ('Manual checks cannot verify PO amount vs GRN', 'Accepted limitation'), ('Vendors ignore calendar', 'Procurement communicates; scorecards')]),
 'C': dict(sheet='Sol C – Supplier Partnership', name='Supplier Quality Partnership', tag='Make vendors and buyers own invoice quality (commercial & governance).',
           eff=S['eff']['S3'], peak=S['peakfix']['S3'], steps=[1, 1, 5, 3, 5, 2, 3, 0, 0], months=4, extra=0,
           capex=[('Vendor master clean-up & onboarding gate', 4), ('Contract clause legal review', 2), ('Vendor segmentation analytics', 2)],
           opex=[('2 supplier quality coordinators', 14)],
           sim=dict(late=S['sol']['S3']['late2'], maxbl=S['sol']['S3']['maxbl']),
           comps=['Vendor segmentation; quarterly supplier quality reviews for top vendors', 'First-time-right certification: guaranteed on-time payment, priority processing, tender score', '"No PO, no pay"; incomplete invoices returned (legal review vs MSME 45-day rule)',
                  'PO before work; GRN posted within 48 h; complete PO lines (Procurement KPI)', 'Vendor master clean-up + onboarding gate (GSTIN, bank)', 'Two supplier quality coordinators run scorecards and coaching'],
           risks=[('Vendor relationships suffer', 'Lead with incentives before consequences'), ('MSME late-payment law', 'Legal review of return-and-restart rule'),
                  ('No vendor-level data today', 'Start vendor coding in month 1'), ('Contract cycles slow', 'Apply to new and renewing contracts first')]),
}
SOLREF = {}
for key, sd in SOLS.items():
    ws = wb.create_sheet(sd['sheet'])
    title(ws, f'Solution {key} · {sd["name"]}', sd['tag'])
    widths(ws, {'A': 50, 'B': 14, 'C': 14, 'D': 14, 'E': 14, 'F': 14, 'G': 4, 'H': 50, 'I': 50})
    # --- KPI block rows 4-12 (filled after we know refs)
    section(ws, 'A4', 'Results (live)')
    # --- error model from row 16
    section(ws, 'A15', '1. Error-reduction model by reason')
    hdr(ws, 16, 1, ['Error reason', 'Baseline share', 'Effectiveness (input)', 'Scenario-adjusted', 'Share of all errors removed'])
    r0 = 17
    for k, rs in enumerate(REASONS):
        row = r0 + k
        put(ws, f'A{row}', rs, sz=9)
        link(ws, f'B{row}', f"='Baseline'!M{5+k}", '0.0%')
        inp(ws, f'C{row}', sd['eff'][k] if isinstance(sd['eff'], list) else list(sd['eff'].values())[k], '0%')
        calc(ws, f'D{row}', f'=MIN(1,C{row}*M_EFF)', '0%')
        calc(ws, f'E{row}', f'=B{row}*D{row}', '0.0%')
    r1 = r0 + len(REASONS) - 1
    rt = r1 + 1
    put(ws, f'A{rt}', 'Reason-level cut (sum)', b=True, fill=F_TOT); calc(ws, f'E{rt}', f'=SUM(E{r0}:E{r1})', '0.0%', True, F_TOT)
    put(ws, f'A{rt+1}', 'Peak-excess errors removed (input)'); inp(ws, f'C{rt+1}', sd['peak'], '0%'); calc(ws, f'D{rt+1}', f'=MIN(1,C{rt+1}*M_EFF)', '0%')
    put(ws, f'A{rt+2}', 'Peak cut = peak-excess share × removal'); link(ws, f'B{rt+2}', '=EXCESS', '0.0%'); calc(ws, f'E{rt+2}', f'=B{rt+2}*D{rt+1}', '0.0%')
    put(ws, f'A{rt+3}', 'TOTAL ERROR CUT = 1 − (1 − reason cut) × (1 − peak cut)', b=True, fill=F_KPI); calc(ws, f'E{rt+3}', f'=1-(1-E{rt})*(1-E{rt+2})', '0.0%', True, F_KPI)
    CUT = f'E{rt+3}'
    ws.conditional_formatting.add(f'E{r0}:E{r1}', DataBarRule(start_type='num', start_value=0, end_type='max', color='5B8BD6'))
    # --- capacity
    c0 = rt + 6
    section(ws, f'A{c0}', '2. Capacity model (analyst touch minutes per invoice)')
    hdr(ws, c0 + 1, 1, ['Step', 'Baseline min', 'Solution min (input)'])
    for k, (stp, bmin, smin) in enumerate(zip(STEPS, BASE_STEPS, sd['steps'])):
        row = c0 + 2 + k
        put(ws, f'A{row}', stp, sz=9); put(ws, f'B{row}', bmin, '0.0', al='center'); inp(ws, f'C{row}', smin, '0.0')
    s1_, s2_ = c0 + 2, c0 + 2 + len(STEPS) - 1
    q = s2_ + 1
    put(ws, f'A{q}', 'Exception share after (25% × (1 − cut))'); calc(ws, f'C{q}', f'=EXC_SHARE*(1-{CUT})', '0.0%')
    put(ws, f'A{q+1}', 'Weighted touch time incl. exceptions (min)', b=True); calc(ws, f'B{q+1}', f'=SUM(B{s1_}:B{s2_})+EXC_SHARE*EXC_MIN', '0.0', True); calc(ws, f'C{q+1}', f'=SUM(C{s1_}:C{s2_})+C{q}*EXC_MIN', '0.0', True, F_KPI)
    put(ws, f'A{q+2}', 'Analyst-equivalents freed on an average day'); calc(ws, f'C{q+2}', f"=(B{q+1}-C{q+1})*'Baseline'!B7/60/8", '0.0')
    put(ws, f'A{q+3}', 'Max invoices per normal day (18 analysts × 8 h)'); calc(ws, f'B{q+3}', f'=18*8*60/B{q+1}', '#,##0'); calc(ws, f'C{q+3}', f'=18*8*60/C{q+1}', '#,##0')
    put(ws, f'A{q+4}', 'Invoices waiting 2+ days (simulation, pasted value)'); put(ws, f'B{q+4}', S['base_des']['late2'], '0%', al='center'); put(ws, f'C{q+4}', sd['sim']['late'], '0%', al='center')
    put(ws, f'A{q+5}', 'Max month-end backlog, invoices (simulation, pasted value)'); put(ws, f'B{q+5}', S['base_des']['maxbl'], '#,##0', al='center'); put(ws, f'C{q+5}', sd['sim']['maxbl'], '#,##0', al='center')
    TOUCH = f'C{q+1}'; FTE = f'C{q+2}'
    # --- costs
    k0 = q + 8
    section(ws, f'A{k0}', '3. Costs (₹ lakh, illustrative – replace with quotes)')
    hdr(ws, k0 + 1, 1, ['Item', 'Base (input)', 'Scenario-adjusted'])
    row = k0 + 2; cap_first = row
    for it, v in sd['capex']:
        put(ws, f'A{row}', 'Capex: ' + it, sz=9); inp(ws, f'B{row}', v, '0.0'); calc(ws, f'C{row}', f'=B{row}*M_COST', '0.0'); row += 1
    cap_last = row - 1
    put(ws, f'A{row}', 'Total capex', b=True, fill=F_TOT); calc(ws, f'C{row}', f'=SUM(C{cap_first}:C{cap_last})', '0.0', True, F_TOT); CAPEX = f'C{row}'; row += 1
    op_first = row
    for it, v in sd['opex']:
        put(ws, f'A{row}', 'Opex/yr: ' + it, sz=9); inp(ws, f'B{row}', v, '0.0'); calc(ws, f'C{row}', f'=B{row}*M_COST', '0.0'); row += 1
    op_last = row - 1
    put(ws, f'A{row}', 'Total opex per year', b=True, fill=F_TOT); calc(ws, f'C{row}', f'=SUM(C{op_first}:C{op_last})', '0.0', True, F_TOT); OPEX = f'C{row}'; row += 2
    # --- benefits
    section(ws, f'A{row}', '4. Benefits (₹ lakh per year)'); row += 1
    put(ws, f'A{row}', 'Baseline COPQ (scenario penalty rate)'); link(ws, f'C{row}', '=COPQ', '0.0'); BCOPQ = f'C{row}'; row += 1
    put(ws, f'A{row}', 'COPQ removed = baseline COPQ × total cut'); calc(ws, f'C{row}', f'={BCOPQ}*{CUT}', '0.0'); BEN1 = f'C{row}'; row += 1
    put(ws, f'A{row}', 'Other operating saving (input; B = staffing avoided by calendar)'); inp(ws, f'B{row}', sd['extra'], '0.0'); calc(ws, f'C{row}', f'=B{row}*M_EFF', '0.0'); BEN2 = f'C{row}'; row += 1
    put(ws, f'A{row}', 'Total benefit per year', b=True, fill=F_TOT); calc(ws, f'C{row}', f'={BEN1}+{BEN2}', '0.0', True, F_TOT); BEN = f'C{row}'; row += 1
    put(ws, f'A{row}', 'COPQ remaining after solution'); calc(ws, f'C{row}', f'={BCOPQ}*(1-{CUT})', '0.0'); COPQA = f'C{row}'; row += 2
    # --- cash flow
    section(ws, f'A{row}', '5. Cash flow & returns (₹ lakh)'); row += 1
    hdr(ws, row, 1, ['Line', 'Year 0', 'Year 1', 'Year 2', 'Year 3']); row += 1
    put(ws, f'A{row}', 'Benefit ramp (input)'); put(ws, f'B{row}', None); [inp(ws, f'{c}{row}', 1.0, '0%') for c in 'CDE']; RAMP = row; row += 1
    put(ws, f'A{row}', 'Capex'); calc(ws, f'B{row}', f'=-{CAPEX}', '0.0'); [calc(ws, f'{c}{row}', 0, '0.0') for c in 'CDE']; CX = row; row += 1
    put(ws, f'A{row}', 'Benefit'); calc(ws, f'B{row}', 0, '0.0'); [calc(ws, f'{c}{row}', f'={BEN}*{c}{RAMP}', '0.0') for c in 'CDE']; BN = row; row += 1
    put(ws, f'A{row}', 'Opex'); calc(ws, f'B{row}', 0, '0.0'); [calc(ws, f'{c}{row}', f'=-{OPEX}', '0.0') for c in 'CDE']; OX = row; row += 1
    put(ws, f'A{row}', 'Net cash flow', b=True, fill=F_TOT); [calc(ws, f'{c}{row}', f'=SUM({c}{CX}:{c}{OX})', '0.0', True, F_TOT) for c in 'BCDE']; NC = row; row += 1
    put(ws, f'A{row}', 'Cumulative'); calc(ws, f'B{row}', f'=B{NC}', '0.0'); [calc(ws, f'{c}{row}', f'={p}{row}+{c}{NC}', '0.0') for p, c in zip('BCD', 'CDE')]; row += 2
    put(ws, f'A{row}', 'Simple NPV (3 yrs, benefits from day 1)', b=True); calc(ws, f'C{row}', f'=B{NC}+NPV(DISC,C{NC}:E{NC})', '0.0', True, F_KPI); NPV_ = f'C{row}'; row += 1
    put(ws, f'A{row}', 'IRR (simple basis, benefits from day 1)', b=True); calc(ws, f'C{row}', f'=IFERROR(IRR(B{NC}:E{NC}),"n/a")', '0%', True); IRR_ = f'C{row}'; row += 1
    put(ws, f'A{row}', 'Simple payback (months, capex ÷ year-1 net)', b=True); calc(ws, f'C{row}', f'=IF(C{NC}>0,{CAPEX}/C{NC}*12,"No payback")', '0.0', True, F_KPI); PB = f'C{row}'; row += 1
    put(ws, f'A{row}', '3-year ROI', b=True); calc(ws, f'C{row}', f'=IFERROR((SUM(C{NC}:E{NC})+B{NC})/{CAPEX},"n/a")', '0%', True); ROI_ = f'C{row}'; row += 1
    put(ws, f'A{row}', 'Time to first impact (months, input)'); inp(ws, f'C{row}', sd['months'], '0.0'); MON = f'C{row}'; row += 1
    put(ws, f'A{row}', 'NPV 3 yrs INCL. BUILD TIME (₹ L) – each part pays from its go-live', b=True); link(ws, f'C{row}', f"='Calc {key}'!B28", '0.0', True); ws[f'C{row}'].fill = F_KPI; NPVB = f'C{row}'; row += 1
    put(ws, f'A{row}', 'Payback from project start (months, incl. build time)', b=True); link(ws, f'C{row}', f"='Calc {key}'!B29", '0', True); ws[f'C{row}'].fill = F_KPI; PBB = f'C{row}'; row += 1
    put(ws, f'A{row}', 'Headline figures count build time (from the hidden Calc sheet, built from the Parts Library). Simple figures above assume benefits from day 1.', sz=8, c=GREY, border=False, i=True); row += 1
    # --- KPI block at top (now refs known)
    kp = [('Total error cut', f'={CUT}', '0.0%'), ('DPMO after', f'=DPMO0*(1-{CUT})', '#,##0'), ('Sigma level after', f'=NORMSINV(1-DPU0*(1-{CUT})/OPP)+1.5', '0.00'),
          ('% invoices with ≥1 error after', f'=1-EXP(-DPU0*(1-{CUT}))', '0.0%'), ('Touch time per invoice (min)', f'={TOUCH}', '0.0'), ('Net benefit per year (₹ L)', f'=C{NC}', '0.0'),
          ('Capex (₹ L)', f'={CAPEX}', '0.0'), ('Payback from start (months)', f'={PBB}', '0'), ('NPV 3 yrs incl. build (₹ L)', f'={NPVB}', '0.0')]
    for k, (lab, fm, nf) in enumerate(kp):
        rr = 5 + k // 3; cc = 'A' if False else None
    # lay KPI as 3 rows × (label,value) pairs in columns A/B, C/D, E/F
    for k, (lab, fm, nf) in enumerate(kp):
        rr = 5 + k % 3 * 2; col = ['A', 'C', 'E'][k // 3]
        lab_cell = ws[f'{col}{rr}']; lab_cell.value = lab; lab_cell.font = f(8, True, GREY); lab_cell.fill = F_KPI
        vcol = col if col != 'A' else 'A'
        vc = ws[f'{col}{rr+1}']; vc.value = fm; vc.number_format = nf; vc.font = f(14, True, NAVY); vc.fill = F_KPI; vc.alignment = Alignment(horizontal='left')
    # components & risks on the right
    section(ws, 'H4', 'Components')
    for k, cmp in enumerate(sd['comps']):
        put(ws, f'H{5+k}', f'{k+1}. {cmp}', wrap=True, sz=9); ws.merge_cells(f'H{5+k}:I{5+k}'); ws.row_dimensions[5 + k].height = 30
    section(ws, 'H13', 'Risks and mitigations'); hdr(ws, 14, 8, ['Risk', 'Mitigation'])
    for k, (rk, mt) in enumerate(sd['risks']):
        put(ws, f'H{15+k}', rk, wrap=True, sz=9); put(ws, f'I{15+k}', mt, wrap=True, sz=9)
    ws.freeze_panes = 'A4'
    SOLREF[key] = dict(sheet=sd['sheet'], CUT=CUT, NPV=NPV_, PB=PB, CAPEX=CAPEX, OPEX=OPEX, BEN=BEN, TOUCH=TOUCH, FTE=FTE, IRR=IRR_, ROI=ROI_, MON=MON, COPQA=COPQA, NC=NC,
                       late=f'C{q+4}', maxbl=f'C{q+5}', r0=r0, peakrow=rt + 1, PBB=PBB, NPVB=NPVB)

# ============================================================ COMBINED
CB = wb.create_sheet('Combined A+B+C')
title(CB, 'Combined: A + B + C', 'Each reason\'s combined effectiveness = 1 − (1 − A)(1 − B)(1 − C). Costs add up (no overlap discount).')
widths(CB, {'A': 50, 'B': 13, 'C': 13, 'D': 13, 'E': 13, 'F': 15, 'G': 15})
hdr(CB, 4, 1, ['Error reason', 'Share', 'A (adj.)', 'B (adj.)', 'C (adj.)', 'Combined effectiveness', 'Share removed'])
for k, rs in enumerate(REASONS):
    row = 5 + k
    put(CB, f'A{row}', rs, sz=9); link(CB, f'B{row}', f"='Baseline'!M{5+k}", '0.0%')
    for col, key in zip('CDE', 'ABC'):
        link(CB, f'{col}{row}', f"='{SOLREF[key]['sheet']}'!D{SOLREF[key]['r0']+k}", '0%')
    calc(CB, f'F{row}', f'=1-(1-C{row})*(1-D{row})*(1-E{row})', '0%'); calc(CB, f'G{row}', f'=B{row}*F{row}', '0.0%')
e = 5 + len(REASONS) - 1
put(CB, f'A{e+1}', 'Reason-level cut', b=True, fill=F_TOT); calc(CB, f'G{e+1}', f'=SUM(G5:G{e})', '0.0%', True, F_TOT)
put(CB, f'A{e+2}', 'Peak cut = excess share × (1 − (1−a)(1−b)(1−c))')
calc(CB, f'G{e+2}', '=EXCESS*(1-' + '*'.join(f"(1-'{SOLREF[k]['sheet']}'!D{SOLREF[k]['peakrow']})" for k in 'ABC') + ')', '0.0%')
put(CB, f'A{e+3}', 'TOTAL COMBINED CUT', b=True, fill=F_KPI); calc(CB, f'G{e+3}', f'=1-(1-G{e+1})*(1-G{e+2})', '0.0%', True, F_KPI)
CCUT = f"'Combined A+B+C'!$G${e+3}"
r = e + 5
lines = [('DPMO after', f'=DPMO0*(1-G{e+3})', '#,##0'), ('Verified DPMO (Improve audit)', "='Baseline'!B20", '#,##0'), ('Verified error reduction', '=VERIFIED', '0.0%'),
         ('Capex (₹ L)', '=' + '+'.join(f"'{SOLREF[k]['sheet']}'!{SOLREF[k]['CAPEX']}" for k in 'ABC'), '0.0'),
         ('Opex per year (₹ L)', '=' + '+'.join(f"'{SOLREF[k]['sheet']}'!{SOLREF[k]['OPEX']}" for k in 'ABC'), '0.0'),
         ('Benefit per year (₹ L)', f"=COPQ*G{e+3}+'{SOLREF['B']['sheet']}'!{SOLREF['B']['BEN']}-'{SOLREF['B']['sheet']}'!{SOLREF['B']['sheet'] and SOLREF['B']['BEN']}+COPQ*0", '0.0')]
# simpler benefit: COPQ × combined cut + B's other saving
lines[-1] = ('Benefit per year (₹ L)', f"=COPQ*G{e+3}+'{SOLREF['B']['sheet']}'!C{int(SOLREF['B']['BEN'][1:])-1}", '0.0')
for k, (lab, fm, nf) in enumerate(lines):
    put(CB, f'A{r+k}', lab, b=True); calc(CB, f'B{r+k}', fm, nf, True)
rb = r
put(CB, f'A{r+6}', 'Net benefit per year (₹ L)', b=True); calc(CB, f'B{r+6}', f'=B{r+5}-B{r+4}', '0.0', True, F_KPI)
put(CB, f'A{r+7}', 'Simple payback (months, benefits from day 1)', b=True); calc(CB, f'B{r+7}', f'=IF(B{r+6}>0,B{r+3}/B{r+6}*12,"No payback")', '0.0', True, F_KPI)
put(CB, f'A{r+8}', 'Simple NPV 3 years (₹ L, benefits from day 1)', b=True); calc(CB, f'B{r+8}', f'=-B{r+3}+B{r+6}/(1+DISC)+B{r+6}/(1+DISC)^2+B{r+6}/(1+DISC)^3', '0.0', True, F_KPI)
put(CB, f'A{r+9}', 'NPV 3 yrs INCL. BUILD TIME (₹ L)', b=True); link(CB, f'B{r+9}', "='Calc A+B+C'!B28", '0.0', True); CB[f'B{r+9}'].fill = F_KPI
put(CB, f'A{r+10}', 'Payback from project start (months, incl. build time)', b=True); link(CB, f'B{r+10}', "='Calc A+B+C'!B29", '0', True); CB[f'B{r+10}'].fill = F_KPI
COMB = dict(npvb=f"'Combined A+B+C'!$B${r+9}", pbb=f"'Combined A+B+C'!$B${r+10}", dpmo=f"'Combined A+B+C'!$B${r}", capex=f"'Combined A+B+C'!$B${r+3}", opex=f"'Combined A+B+C'!$B${r+4}", ben=f"'Combined A+B+C'!$B${r+5}",
            net=f"'Combined A+B+C'!$B${r+6}", pb=f"'Combined A+B+C'!$B${r+7}", npv=f"'Combined A+B+C'!$B${r+8}")
CB.freeze_panes = 'A5'

# ============================================================ COMPARISON
CP = wb.create_sheet('Comparison')
title(CP, 'Comparison & decision matrix', 'All values link to the solution sheets. Change weights or judgement scores (yellow) to test the ranking.')
widths(CP, {'A': 40, 'B': 17, 'C': 17, 'D': 17, 'E': 17, 'F': 17, 'G': 40})
hdr(CP, 4, 1, ['Metric', 'A Digital Front Door', 'B Lean Flow', 'C Supplier Partnership', 'A + B + C', 'Hybrid H'])
def ref(k, key): return f"='{SOLREF[k]['sheet']}'!{SOLREF[k][key]}"
cmp_rows = [('Total error cut', 'CUT', '0.0%', f'={CCUT}'), ('DPMO after', None, '#,##0', f'={COMB["dpmo"]}'),
            ('Touch time per invoice (min)', 'TOUCH', '0.0', ''), ('Analyst-equivalents freed', 'FTE', '0.0', ''),
            ('Invoices waiting 2+ days (sim.)', 'late', '0%', ''), ('Capex (₹ L)', 'CAPEX', '0.0', f'={COMB["capex"]}'),
            ('Opex per year (₹ L)', 'OPEX', '0.0', f'={COMB["opex"]}'), ('Benefit per year (₹ L)', 'BEN', '0.0', f'={COMB["ben"]}'),
            ('Payback from project start (months)', 'PBB', '0', f'={COMB["pbb"]}'), ('NPV 3 yrs incl. build time (₹ L)', 'NPVB', '0.0', f'={COMB["npvb"]}'), ('IRR (simple basis)', 'IRR', '0%', ''),
            ('Time to first impact (months)', 'MON', '0.0', '')]
for k, (lab, key, nf, comb) in enumerate(cmp_rows):
    row = 5 + k
    put(CP, f'A{row}', lab, b=True)
    for col, s in zip('BCD', 'ABC'):
        if key: link(CP, f'{col}{row}', ref(s, key), nf)
        else: link(CP, f'{col}{row}', f"=DPMO0*(1-'{SOLREF[s]['sheet']}'!{SOLREF[s]['CUT']})", nf)
    if comb: link(CP, f'E{row}', comb, nf)
    else: put(CP, f'E{row}', '–', al='center')
put(CP, 'A17', 'Verified result (Improve audit)', b=True); link(CP, 'B17', '=VERIFIED', '0.0%'); put(CP, 'C17', 'error cut, Mar–Apr 2027 vs adjusted baseline', sz=9, c=GREY, border=False)
# decision matrix
section(CP, 'A19', 'Weighted decision matrix')
hdr(CP, 20, 1, ['Criterion', 'Weight (input)', 'A', 'B', 'C'])
crit = [('Impact on errors (cut × 10)', 0.35, [f'=B5*10', f'=C5*10', f'=D5*10']),
        ('Cost: 10 × (1 − 3-yr cost ÷ ₹100 L)', 0.20, [f'=MAX(0,10*(1-(B10+3*B11)/100))', f'=MAX(0,10*(1-(C10+3*C11)/100))', f'=MAX(0,10*(1-(D10+3*D11)/100))']),
        ('Speed: 10 − 1.2 × months to impact', 0.20, [f'=10-1.2*B16', f'=10-1.2*C16', f'=10-1.2*D16']),
        ('Risk, higher = safer (judgement)', 0.15, [6, 8, 5]), ('Scalability & sustainability (judgement)', 0.10, [9, 4, 7])]
for k, (lab, w, sc) in enumerate(crit):
    row = 21 + k
    put(CP, f'A{row}', lab); inp(CP, f'B{row}', w, '0%')
    for col, v in zip('CDE', sc):
        (calc(CP, f'{col}{row}', v, '0.0') if isinstance(v, str) else inp(CP, f'{col}{row}', v, '0.0'))
put(CP, 'A26', 'Weighted score', b=True, fill=F_TOT); calc(CP, 'B26', '=SUM(B21:B25)', '0%', True, F_TOT)
for col in 'CDE': calc(CP, f'{col}26', f'=SUMPRODUCT($B$21:$B$25,{col}21:{col}25)', '0.00', True, F_KPI)
put(CP, 'A27', 'Rank'); [calc(CP, f'{col}27', f'=RANK({col}26,$C$26:$E$26)', '0') for col in 'CDE']
put(CP, 'A28', 'Meets ≤ 500 DPMO stretch alone? (team-set, not in the brief)'); [calc(CP, f'{col}28', f'=IF({b}6<=500,"Yes","No")', None) for col, b in zip('CDE', 'BCD')]
put(CP, 'A29', 'Clears the ≥ 75% cut pass mark alone?'); [calc(CP, f'{col}29', f'=IF({b}5>=0.75,"Yes","No")', None) for col, b in zip('CDE', 'BCD')]
put(CP, 'A30', 'Reading: the matrix favours the cheapest, fastest option. Only A clears the ≥ 75% pass mark alone, narrowly (needs ≥ 89% of its assumed effect) and late (~5 months); none reaches the ≤ 500 DPMO stretch. The path is B first, build A, embed C\'s PO/GRN discipline – refined by the stress test into Hybrid H. The A+B+C model (−95%) is consistent with, not proof of, the supplied audit (−96%).', wrap=True, sz=9)
CP.merge_cells('A30:E30'); CP.row_dimensions[30].height = 42
CP.conditional_formatting.add('C27:E27', CellIsRule(operator='equal', formula=['1'], fill=PatternFill('solid', fgColor='C6EFCE'), font=Font(name=FONT, bold=True, color='006100')))
# chart data
section(CP, 'G4', 'Chart data')
hdr(CP, 5, 7, ['Option', 'Error cut %', 'NPV ₹ L', 'Payback months'])
CP.column_dimensions['H'].width = 13; CP.column_dimensions['I'].width = 13; CP.column_dimensions['J'].width = 15
opts = [('A Digital Front Door', 'B'), ('B Lean Flow', 'C'), ('C Supplier Partnership', 'D'), ('A + B + C', 'E'), ('Hybrid H', 'F')]
for k, (lab, col) in enumerate(opts):
    row = 6 + k
    put(CP, f'G{row}', lab); calc(CP, f'H{row}', f'={col}5*100', '0'); calc(CP, f'I{row}', f'={col}14', '0'); calc(CP, f'J{row}', f'={col}13', '0.0')
put(CP, 'G11', 'Verified (Improve audit)'); calc(CP, 'H11', '=VERIFIED*100', '0')
def bar(t, col, rows_, anchor, color):
    ch = BarChart(); ch.type = 'bar'; ch.gapWidth = 50; ch.title = t; ch.legend = None; ch.width = 15; ch.height = 7
    ch.add_data(Reference(CP, min_col=col, min_row=6, max_row=5 + rows_), titles_from_data=False)
    ch.set_categories(Reference(CP, min_col=7, min_row=6, max_row=5 + rows_))
    ch.series[0].graphicalProperties.solidFill = color; ch.series[0].graphicalProperties.line.solidFill = color
    ch.dataLabels = DataLabelList(); ch.dataLabels.showVal = True; ch.dataLabels.showSerName = False; ch.dataLabels.showCatName = False; ch.dataLabels.showLegendKey = False
    ch.x_axis.delete = False; ch.y_axis.delete = False; ch.x_axis.scaling.orientation = 'maxMin'
    CP.add_chart(ch, anchor)
bar('Error cut (%)', 8, 6, 'G13', '13294B'); bar('3-year NPV incl. build time (₹ lakh)', 9, 5, 'G28', '2E5C8A'); bar('Payback from project start (months)', 10, 5, 'G43', 'E8833A')
CP.freeze_panes = 'A5'

# ============================================================ SENSITIVITY
SN = wb.create_sheet('Sensitivity')
title(SN, 'Sensitivity', 'Net benefit per year (₹ L) when one driver changes; everything else at the active scenario. Cut is scaled linearly with the effectiveness multiplier (capped at 97%).')
widths(SN, {'A': 30, 'B': 14, 'C': 14, 'D': 14, 'E': 14})
section(SN, 'A4', '1. Penalty trigger rate')
hdr(SN, 5, 1, ['Penalty trigger rate', 'A', 'B', 'C'])
for k, pr in enumerate([0.01, 0.02, 0.03, 0.05, 0.08, 0.10]):
    row = 6 + k; inp(SN, f'A{row}', pr, '0%')
    for col, s in zip('BCD', 'ABC'):
        sr = SOLREF[s]; sh = f"'{sr['sheet']}'"
        other = f"({sh}!{sr['BEN']}-COPQ*{sh}!{sr['CUT']})"
        calc(SN, f'{col}{row}', f'=(VOL*MONTHS*DEF_SHARE*A{row}*PEN_AMT/100000+COPQ_RW)*{sh}!{sr["CUT"]}+{other}-{sh}!{sr["OPEX"]}', '0.0')
section(SN, 'A14', '2. Effectiveness multiplier (relative to the active scenario)')
hdr(SN, 15, 1, ['Multiplier', 'A', 'B', 'C'])
for k, m in enumerate([0.4, 0.6, 0.8, 1.0, 1.2]):
    row = 16 + k; inp(SN, f'A{row}', m, '0.0')
    for col, s in zip('BCD', 'ABC'):
        sr = SOLREF[s]; sh = f"'{sr['sheet']}'"
        other = f"({sh}!{sr['BEN']}-COPQ*{sh}!{sr['CUT']})"
        calc(SN, f'{col}{row}', f'=COPQ*MIN(0.97,{sh}!{sr["CUT"]}*A{row})+{other}*A{row}-{sh}!{sr["OPEX"]}', '0.0')
section(SN, 'A23', '3. Capex overrun – simple payback months (benefits from day 1)')
hdr(SN, 24, 1, ['Capex multiplier', 'A', 'B', 'C'])
for k, m in enumerate([1.0, 1.5, 2.0, 3.0]):
    row = 25 + k; inp(SN, f'A{row}', m, '0.0')
    for col, s in zip('BCD', 'ABC'):
        sr = SOLREF[s]; sh = f"'{sr['sheet']}'"
        calc(SN, f'{col}{row}', f'=IF({sh}!C{sr["NC"]}>0,{sh}!{sr["CAPEX"]}*A{row}/{sh}!C{sr["NC"]}*12,"No payback")', '0.0')
for rg in ('B6:D11', 'B16:D20'):
    SN.conditional_formatting.add(rg, CellIsRule(operator='lessThan', formula=['0'], font=Font(name=FONT, color='C00000', bold=True)))

# ============================================================ ASSUMPTIONS
AS = wb.create_sheet('Assumptions')
title(AS, 'Assumptions log', 'Every assumption behind the model. IDs match the website\'s assumptions register where relevant.')
widths(AS, {'A': 8, 'B': 90, 'C': 50})
hdr(AS, 4, 1, ['ID', 'Assumption', 'Where used'])
al = [('A3', '25-Aug-2026 (0 errors on a 659-invoice peak day) excluded from the adjusted baseline as a suspected recording anomaly (not confirmable from the data); the raw baseline is kept as a check (verified cut 95.7% vs 95.9%).', 'Baseline DPU, DPMO'),
      ('A4/A5', '58 equal opportunities per invoice; sigma = NORMSINV(1 − DPO) + 1.5.', 'DPMO, sigma'),
      ('A6', 'Share of invoices with ≥1 error = 1 − e^(−DPU) (Poisson).', 'KPIs'),
      ('A9', 'Exception share falls in proportion to the error cut.', 'Capacity model'),
      ('A10', 'Monthly volume = baseline average per day × 26, held constant for all options.', 'COPQ, benefits'),
      ('A11', 'Case cost values: ₹5,000 penalty, 5% trigger, 35% invoices with errors, 25% exceptions, 12 min, ₹350/h.', 'COPQ'),
      ('A12', 'Capex and opex lines are illustrative placeholders; replace with quotes.', 'Costs'),
      ('A13+', 'Effectiveness per reason per solution is analyst judgement (see each solution sheet).', 'Error model'),
      ('P1', 'Peak-excess errors = errors on peak days above the non-peak rate (≈21% of all errors); each solution removes a share of them.', 'Error model'),
      ('P2', 'Headline NPV and payback count build time (each part pays from the month after go-live; hidden Calc sheets). Simple NPV/payback assume benefits from day 1. 3-year horizon; 12% discount rate.', 'Cash flow'),
      ('P3', 'Solution B\'s extra saving = staffing cost avoided by the month-end calendar vs meeting the service level by staffing alone (LP model).', 'Sol B benefit'),
      ('P4', 'Combined option: effectiveness combines as 1 − (1 − a)(1 − b)(1 − c); costs simply add.', 'Combined'),
      ('P5', 'Simulation outputs (waiting 2+ days, max backlog) are pasted values from a Python discrete-event model (20 runs).', 'Capacity'),
      ('P6', 'Scenario multipliers: Conservative = effectiveness × 0.6, costs × 1.5, 2% penalty rate; Optimistic = × 1.15, × 0.8, 8%.', 'All')]
for k, (a, b, c) in enumerate(al):
    row = 5 + k; put(AS, f'A{row}', a, b=True, al='center'); put(AS, f'B{row}', b, wrap=True); put(AS, f'C{row}', c, sz=9); AS.row_dimensions[row].height = 30

exec(open(paths.WORKBOOK / 'ext_hybrid.py').read())
exec(open(paths.WORKBOOK / 'ext_uncert.py').read())
exec(open(paths.WORKBOOK / 'ext_control.py').read())
exec(open(paths.WORKBOOK / 'ext_recon.py').read())
nr = RD.max_row + 1
for a_, b_ in [('Uncertainty & Decision', 'How sure we are: exact Poisson, quasi-Poisson, day-level bootstrap and negative-binomial intervals for the before/after DPMO and the cut; P(true DPMO > 500); Monte Carlo intervals and standard errors for Hybrid H; the end-October (G1) decision rule with thresholds, audit precision, zone outcomes and policy comparison; other statistics quoted in the deck.'),
               ('Control Chart', 'Standard u-chart vs Laney u′ (dispersion-adjusted) limits for the baseline and the supplied audit, all live. With adjusted limits no baseline day (25-Aug excluded) is a special cause; the after-period has σz < 1, so standard limits apply.'),
               ('Four kinds of number', '1) Observed: case data (Baseline, Measure/Improve daily rows). 2) Calculated: base-case model results, live formulas you can change through yellow inputs. 3) Pasted simulation: Monte Carlo, queue simulation, combination search, bootstrap and regression outputs from the Python model (grey cells); change them by rerunning the model. 4) Sensitivity: results under assumed ranges or single-factor stresses (Stress Test, Sensitivity).'),
               ('Deck Reconciliation', 'Every number on the limit_3 deck mapped to a live workbook cell, rounded the way the deck prints it, with a ✓/✗ check. Cell B4 counts the matches.')]:
    put(RD, f'A{nr}', a_, b=True, fill=F_SUB); put(RD, f'B{nr}', b_, wrap=True); RD.row_dimensions[nr].height = 44; nr += 1
na = AS.max_row + 1
for a_, b_, c_ in [('H12', 'Decision rule: c₁ = 1 − DPU(Oct) ÷ baseline DPU, measured on a 20% audit in October. Stay with H if c₁ ≥ 40% (effectiveness ≥ 0.75 of plan); escalate to H+ if 34% ≤ c₁ < 40%; re-diagnose below 34%. Thresholds come from the ramp model and are team proposals.', 'Uncertainty & Decision'),
                   ('H13', 'Uncertainty: baseline errors are over-dispersed (φ ≈ 4.4), so day-level bootstrap or negative-binomial intervals are used for the baseline and the cut; the audit (φ ≈ 0.74) is close to Poisson, so its exact Poisson interval stands.', 'Uncertainty & Decision')]:
    put(AS, f'A{na}', a_, b=True, al='center'); put(AS, f'B{na}', b_, wrap=True); put(AS, f'C{na}', c_, sz=9); AS.row_dimensions[na].height = 42; na += 1
UD.sheet_properties.tabColor = '7030A0'; RC.sheet_properties.tabColor = '7030A0'
wb.calculation.fullCalcOnLoad = True
wb.save(paths.out('model2.xlsx')); print('saved')

# ---- packaging glue (no model logic): optional LibreOffice step that stores calculated values in the file --------------
# Excel recalculates model2.xlsx on opening (fullCalcOnLoad), so this is only a convenience. Skip it with:  --no-recalc
if '--no-recalc' not in sys.argv:
    import subprocess
    subprocess.run([sys.executable, str(paths.WORKBOOK / 'recalc_and_inject.py')], check=False)
