"""Stage 5 (OPTIONAL) - copy computed values from a recalculated workbook into the formulas-only workbook.
Usage : python inject.py <formulas.xlsx> <recalculated.xlsx> <output.xlsx>
Reads : the first two files given on the command line
Writes: the third file (same as the first, but every formula cell also carries its cached value)
"""
# ---------------------------------------------------------------------------
# WHAT THIS FILE DOES (plain English)
#   openpyxl cannot calculate formulas, and LibreOffice (which can) rewrites styles. So this script takes the numbers LibreOffice
#   calculated and writes them into the XML of the original file, keeping its formatting. Only needed if you want a file that shows
#   numbers in viewers that do not calculate (previews, pandas, openpyxl data_only). Excel itself does not need it.
# ---------------------------------------------------------------------------
import zipfile, re, sys, html
from openpyxl import load_workbook
src, calc, out = sys.argv[1:4]
vals = load_workbook(calc, data_only=True)
zin = zipfile.ZipFile(src)
wbx = zin.read('xl/workbook.xml').decode(); rels = zin.read('xl/_rels/workbook.xml.rels').decode()
rid2t = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="/?(?:xl/)?([^"]+)"', rels)) | {a: b for b, a in re.findall(r'Target="/?(?:xl/)?([^"]+)"[^>]*Id="(rId\d+)"', rels)}
sheets = {('xl/' + rid2t[rid]): html.unescape(name) for name, rid in re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="(rId\d+)"', wbx)}
pat = re.compile(r'<c r="([A-Z]+\d+)"([^>]*)>(<f[^>]*>.*?</f>|<f[^>]*/>)<v>[^<]*</v></c>', re.S)
n = 0
zout = zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename in sheets:
        ws = vals[sheets[item.filename]]
        def rep(m):
            global n
            ref, attrs, f = m.groups(); v = ws[ref].value
            attrs = re.sub(r'\s+t="[^"]*"', '', attrs)
            if v is None: return f'<c r="{ref}"{attrs}>{f}</c>'
            n += 1
            if isinstance(v, bool): return f'<c r="{ref}"{attrs} t="b">{f}<v>{int(v)}</v></c>'
            if isinstance(v, (int, float)): return f'<c r="{ref}"{attrs}>{f}<v>{repr(float(v)) if isinstance(v, float) else v}</v></c>'
            s = str(v)
            if s.startswith('#'): return f'<c r="{ref}"{attrs} t="e">{f}<v>{html.escape(s)}</v></c>'
            return f'<c r="{ref}"{attrs} t="str">{f}<v>{html.escape(s, quote=False)}</v></c>'
        data = pat.sub(rep, data.decode()).encode()
    zout.writestr(item, data)
zout.close(); print('injected', n)
